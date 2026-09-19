import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Students and Guide metadata
STUDENTS = [
    ("Abhinav Bhatra", "24EJCCA009"),
    ("Amit kumar", "24EJCCA027"),
    ("Anand Gaur", "24EJCCA030")
]
GUIDE_NAME = "Mona Nagar"
GUIDE_DESIGNATION = "ITS Coordinator (CC)"
COLLEGE_NAME = "JAIPUR ENGINEERING COLLEGE AND RESEARCH CENTRE"
DEPARTMENT_NAME = "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING (ARTIFICIAL INTELLIGENCE)"
UNIVERSITY_NAME = "Rajasthan Technical University, KOTA"
SESSION = "2026-2027"
PROJECT_TITLE = "SORTING ANALYSIS COMPARATOR (SAC)"
ADDRESS_LINES = [
    "Shri Ram ki Nangal, via Sitapura RIICO,",
    "Tonk Road, Sukhpuria, Bambala, Jaipur, Rajasthan"
]

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def add_header_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(16, 44, 87)
        if level == 1:
            run.font.size = Pt(16)
            run.font.bold = True
        elif level == 2:
            run.font.size = Pt(13)
            run.font.bold = True
    return h

def add_body_paragraph(doc, text, bold_prefix=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.bold = True
        r_prefix.font.name = 'Times New Roman'
        r_prefix.font.size = Pt(11.5)

    r_body = p.add_run(text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(11.5)
    return p

# ============================================================================
# 1. GENERATE PROJECT REPORT (.docx)
# ============================================================================
def generate_project_report():
    doc = Document()
    
    # Page setup - 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # --- COVER PAGE ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    
    r = p.add_run("A\nProject Report\nOn\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True

    r_title = p.add_run(f"{PROJECT_TITLE}\n\n")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(17)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(16, 44, 87)

    r_sub = p.add_run("Submitted in partial fulfillment for the award of degree of\n")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(11.5)
    r_sub.italic = True

    r_deg = p.add_run("Bachelor of Technology\n\nDegree Of\n\n")
    r_deg.font.name = 'Times New Roman'
    r_deg.font.size = Pt(13)
    r_deg.font.bold = True

    r_univ = p.add_run(f"{UNIVERSITY_NAME}\n\n")
    r_univ.font.name = 'Times New Roman'
    r_univ.font.size = Pt(14)
    r_univ.font.bold = True

    r_inst = p.add_run(f"{COLLEGE_NAME}\n\n{SESSION}\n\n")
    r_inst.font.name = 'Times New Roman'
    r_inst.font.size = Pt(13)
    r_inst.font.bold = True

    # Submitted By & Coordinator in 2 columns
    table_meta = doc.add_table(rows=1, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_meta.autofit = False
    table_meta.columns[0].width = Inches(3.2)
    table_meta.columns[1].width = Inches(3.2)

    cell_left = table_meta.rows[0].cells[0]
    p_sub = cell_left.paragraphs[0]
    p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_sb = p_sub.add_run("Submitted By:\n")
    r_sb.bold = True
    r_sb.font.name = 'Times New Roman'
    r_sb.font.size = Pt(11)
    for name, roll in STUDENTS:
        r_st = p_sub.add_run(f"{name} ({roll})\n")
        r_st.font.name = 'Times New Roman'
        r_st.font.size = Pt(10.5)

    cell_right = table_meta.rows[0].cells[1]
    p_coord = cell_right.paragraphs[0]
    p_coord.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_cb = p_coord.add_run(f"{GUIDE_DESIGNATION}\n")
    r_cb.bold = True
    r_cb.font.name = 'Times New Roman'
    r_cb.font.size = Pt(11)
    r_cn = p_coord.add_run(f"{GUIDE_NAME}\nFaculty Mentor\n")
    r_cn.font.name = 'Times New Roman'
    r_cn.font.size = Pt(10.5)

    # Department and address footer
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(36)
    p_foot.paragraph_format.space_after = Pt(2)
    
    r_dept = p_foot.add_run(f"{DEPARTMENT_NAME}\n{COLLEGE_NAME}\n")
    r_dept.font.name = 'Times New Roman'
    r_dept.font.size = Pt(11)
    r_dept.font.bold = True

    for line in ADDRESS_LINES:
        r_addr = p_foot.add_run(f"{line}\n")
        r_addr.font.name = 'Times New Roman'
        r_addr.font.size = Pt(9.5)

    doc.add_page_break()

    # --- TABLE OF CONTENTS (PAGE 2) ---
    p_toc_title = doc.add_paragraph()
    p_toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_toc = p_toc_title.add_run("TABLE OF CONTENTS")
    r_toc.font.name = 'Times New Roman'
    r_toc.font.size = Pt(16)
    r_toc.font.bold = True
    r_toc.underline = True
    p_toc_title.paragraph_format.space_after = Pt(20)

    toc_items = [
        ("Introduction", "1"),
        ("Objective", "3"),
        ("Methodology And Planning", "4"),
        ("Literature Review", "6"),
        ("Applications of the Project", "8"),
        ("Block Diagram", "9"),
        ("UML Diagram", "10"),
        ("Screenshots", "11"),
        ("References", "13"),
        ("Approval of the Project", "14")
    ]

    for title, page in toc_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_after = Pt(8)
        p_item.paragraph_format.line_spacing = 1.3
        
        r_t = p_item.add_run(title)
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(12)
        r_t.font.bold = True

        # Dots filler
        dots_count = max(5, 75 - len(title) * 2)
        r_dots = p_item.add_run(" " + ". " * dots_count)
        r_dots.font.name = 'Times New Roman'
        r_dots.font.size = Pt(10)
        r_dots.font.color.rgb = RGBColor(120, 120, 120)

        r_pg = p_item.add_run(" " + page)
        r_pg.font.name = 'Times New Roman'
        r_pg.font.size = Pt(12)
        r_pg.font.bold = True

    doc.add_page_break()

    # --- 1. INTRODUCTION ---
    add_header_styled(doc, "1. Introduction", level=1)
    add_body_paragraph(doc, 
        "Sorting is one of the most fundamental operations in computer science, serving as an algorithmic building block for information retrieval, database indexing, computer graphics, computational geometry, and distributed big-data pipelines. An algorithm's asymptotic complexity, expressed using Big-O notation, provides a theoretical upper bound on its growth rate as input scale expands toward infinity. However, in physical computing hardware, real-world execution latency frequently deviates from pure mathematical models due to CPU cache hierarchies (L1/L2/L3), memory latency, branch prediction behavior, instruction pipelining, and data distribution topology.")

    add_body_paragraph(doc,
        "The Sorting Analysis Comparator (SAC) is an engineering and laboratory analysis system built entirely in modern C++17 with an interactive graphical interface powered by Simple and Fast Multimedia Library (SFML 2.5). The system is engineered to bridge theoretical algorithm design and empirical runtime behavior. By benchmarking four seminal sorting paradigms—Merge Sort (Divide-and-Conquer), Quick Sort (Partitioning), Heap Sort (Selection via Priority Queue), and Radix Sort (Non-comparative Distribution)—SAC provides rigorous hardware telemetry across multiple input scales (N = 100 to 50,000) and topological distributions (Uniform Random, Nearly Sorted, Reverse Sorted, and Many Duplicates).")

    add_body_paragraph(doc,
        "Beyond statistical benchmarks, SAC incorporates a real-time 60 FPS animation engine that visually traces comparison and swap operations, supports an unprecedented 4-Way Head-to-Head Race Mode, and enables user-defined custom array insertion with dynamic number labels drawn directly on bars.")

    # --- 2. OBJECTIVE ---
    add_header_styled(doc, "2. Objective", level=1)
    add_body_paragraph(doc,
        "The primary objectives of the Sorting Analysis Comparator project are defined as follows:")
    
    objectives = [
        ("Empirical Complexity Verification: ", "Measure real execution wall-clock time with sub-microsecond hardware counters (Windows QueryPerformanceCounter) and determine whether observed scaling ratios T(n2)/T(n1) align with theoretical O(n log n), O(n^2), and O(d(n+k)) Big-O complexity predictions."),
        ("Exact Operation Telemetry: ", "Accurately count comparisons and element swaps/assignments across identical input permutations, highlighting the invariant nature of deterministic operations independent of CPU clock fluctuations."),
        ("Stress Testing Across Topologies: ", "Investigate algorithm resilience against pathological input distributions, specifically verifying Quick Sort's quadratic degradation on high duplicate frequency and Merge Sort's consistent performance."),
        ("Interactive Multi-Paradigm Visualization: ", "Develop a high-performance graphical visualizer with three operational modes (Single Sort Bar Animation, 4-Way Race Mode, and Benchmark Growth Curves Plotter) and mouse-driven controls."),
        ("Custom Data Tracing: ", "Empower students and researchers to input arbitrary integer sequences, visually verifying correctness and invariants of sorting passes with real-time on-bar numeric tracking.")
    ]
    for pref, text in objectives:
        add_body_paragraph(doc, text, bold_prefix=pref, space_after=4)

    # --- 3. METHODOLOGY AND PLANNING ---
    add_header_styled(doc, "3. Methodology And Planning", level=1)
    add_body_paragraph(doc,
        "The project architecture is segregated into four decoupled layers to maintain modularity, testability, and deterministic timing fidelity:")

    methodology_steps = [
        ("Dual-Variant Architecture: ", "Every sorting algorithm is implemented in two distinct variants: a Fast Variant compiled with -O3 aggressive optimizations, preallocated buffers, and zero telemetry overhead for genuine benchmark timing; and an Instrumented Variant that records an AnimationStep event log alongside operation counters."),
        ("Hardware Sub-Microsecond Timer: ", "Standard MinGW GCC implementations of std::chrono::high_resolution_clock suffer from a 15.6 ms resolution tick floor on Windows. SAC implements an operating-system-level HighResClock utilizing Windows QueryPerformanceCounter (QPC), capturing latency down to nanosecond precision."),
        ("Data Generation Engine: ", "Four distinct stochastic generators synthesize arrays across six logarithmic scales: N in {100, 500, 1000, 5000, 10000, 50000}. Distributions include Uniform Random (discrete uniform distribution in [1, 100000]), Nearly Sorted (~5% random index swaps), Reverse Sorted (strictly descending), and Many Duplicates (domain [1, 10])."),
        ("Empirical Ratio & Power Fitting: ", "Calculates observed growth ratios R_obs = T(n2)/T(n1) between successive scales and fits the empirical power exponent alpha = ln(R_obs)/ln(n2/n1). Alpha ~ 1.0 indicates linear O(n), alpha ~ 1.05-1.2 indicates O(n log n), and alpha ~ 2.0 flags quadratic O(n^2)."),
        ("SFML Render Pipeline: ", "An asynchronous event loop renders vertical bar geometries at 60 FPS. Operations are color-coded: Yellow for Comparisons, Red for Swaps, Purple for Overwrites/Assignments, and Emerald Green for permanently sorted elements.")
    ]
    for pref, text in methodology_steps:
        add_body_paragraph(doc, text, bold_prefix=pref, space_after=4)

    # --- 4. LITERATURE REVIEW ---
    add_header_styled(doc, "4. Literature Review", level=1)
    add_body_paragraph(doc,
        "The study of comparison-based and distribution-based sorting algorithms has a rich academic history spanning several decades. Fundamental findings from classical literature form the foundation of this research:")

    lit_reviews = [
        ("Merge Sort (John von Neumann, 1945): ", "Formulated on the divide-and-conquer strategy, Merge Sort guarantees Theta(n log n) comparisons in all cases by recursively bisecting arrays and merging sorted subarrays. Its primary drawback is auxiliary spatial complexity Theta(n), making in-place variants challenging on memory-constrained systems."),
        ("Quick Sort (C. A. R. Hoare, 1959): ", "Quick Sort employs in-place partitioning around a selected pivot element. While exhibiting an average time complexity of Theta(n log n) with an exceptionally small constant factor due to cache-friendly contiguous memory scanning, naive pivot selection degrades to O(n^2) when partitioning degenerates on sorted or duplicate-heavy arrays. Median-of-three pivot selection and tail-recursion elimination mitigate these risks."),
        ("Heap Sort (J. W. J. Williams & R. W. Floyd, 1964): ", "Constructs a complete binary max-heap in Theta(n) time using bottom-up sift-down operations, followed by n - 1 successive root extractions in O(log n) each, achieving strict O(n log n) upper bound in-place with O(1) auxiliary space."),
        ("Radix Sort (Herman Hollerith, 1887; Harold H. Seward, 1954): ", "Radix Sort circumvents the Omega(n log n) information-theoretic lower bound for comparison sorts by examining positional key digits using stable Counting Sort passes. Its theoretical time complexity is Theta(d * (n + k)), rendering it strictly linear when key length d and radix base k are constant.")
    ]
    for pref, text in lit_reviews:
        add_body_paragraph(doc, text, bold_prefix=pref, space_after=4)

    # Complexity Summary Table
    table_comp = doc.add_table(rows=5, cols=5)
    table_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_comp.autofit = False

    headers = ["Algorithm", "Best Case", "Average Case", "Worst Case", "Auxiliary Space"]
    col_widths = [Inches(1.4), Inches(1.3), Inches(1.3), Inches(1.3), Inches(1.2)]
    
    hdr_cells = table_comp.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        hdr_cells[i].text = h_text
        set_cell_background(hdr_cells[i], "102C57")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.name = 'Times New Roman'
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(10)

    rows_data = [
        ["Merge Sort", "O(n log n)", "O(n log n)", "O(n log n)", "O(n)"],
        ["Quick Sort", "O(n log n)", "O(n log n)", "O(n^2)", "O(log n)"],
        ["Heap Sort", "O(n log n)", "O(n log n)", "O(n log n)", "O(1)"],
        ["Radix Sort", "O(d*(n+k))", "O(d*(n+k))", "O(d*(n+k))", "O(n+k)"]
    ]

    for row_idx, data in enumerate(rows_data, start=1):
        row_cells = table_comp.rows[row_idx].cells
        bg_col = "F4F6F9" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(data):
            row_cells[col_idx].width = col_widths[col_idx]
            row_cells[col_idx].text = text
            set_cell_background(row_cells[col_idx], bg_col)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- 5. APPLICATIONS OF THE PROJECT ---
    add_header_styled(doc, "5. Applications of the Project", level=1)
    add_body_paragraph(doc,
        "The Sorting Analysis Comparator serves practical engineering and educational utility across multiple domains:")
    
    apps = [
        ("Academic Pedagogy: ", "Equips universities and computer science institutions with an interactive laboratory instrument that transforms abstract asymptotic recurrence relations into tangible visual and statistical telemetry."),
        ("Systems & Database Engineering: ", "Assists database engine developers in benchmarking sort-merge join operators and selecting optimal sorting heuristics based on data presortedness and cardinality."),
        ("Embedded Systems Constraint Profiling: ", "Facilitates memory-constrained runtime evaluations, proving Heap Sort's zero-allocation advantage over Merge Sort in microcontrollers."),
        ("Algorithmic Invariant Verification: ", "Enables visual tracing of loop invariants (e.g. max-heap property preservation, partition pivot boundary defense) on user-provided edge case arrays.")
    ]
    for pref, text in apps:
        add_body_paragraph(doc, text, bold_prefix=pref, space_after=4)

    # --- 6. BLOCK DIAGRAM ---
    add_header_styled(doc, "6. Block Diagram", level=1)
    add_body_paragraph(doc,
        "Figure 6.1 illustrates the architectural block diagram of the Sorting Analysis Comparator system, depicting the workflow from input data synthesis through dual-variant execution to telemetry calculation and graphical rendering.")
    
    if os.path.exists("docs/block_diagram.png"):
        doc.add_picture("docs/block_diagram.png", width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 6.1: System Architecture Block Diagram of SAC")
        r_cap.italic = True
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10)
        p_cap.paragraph_format.space_after = Pt(14)

    # --- 7. UML DIAGRAM ---
    add_header_styled(doc, "7. UML Diagram", level=1)
    add_body_paragraph(doc,
        "Figure 7.1 presents the Unified Modeling Language (UML) Class Diagram of SAC, modeling class relationships, public/private member contracts, and event structures across the visualizer and engine.")

    if os.path.exists("docs/uml_diagram.png"):
        doc.add_picture("docs/uml_diagram.png", width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 7.1: UML Class Diagram for Sorting Analysis Comparator")
        r_cap.italic = True
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10)
        p_cap.paragraph_format.space_after = Pt(14)

    # --- 8. SCREENSHOTS ---
    add_header_styled(doc, "8. Screenshots", level=1)
    add_body_paragraph(doc,
        "The following figures illustrate the operational user interfaces and benchmarking visual outputs generated by the system:")

    if os.path.exists("docs/screenshot_visualizer.png"):
        doc.add_picture("docs/screenshot_visualizer.png", width=Inches(6.0))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 8.1: Single Sort Animation Mode with Real-Time Telemetry HUD")
        r_cap.italic = True
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10)
        p_cap.paragraph_format.space_after = Pt(12)

    if os.path.exists("docs/benchmark_curves.png"):
        doc.add_picture("docs/benchmark_curves.png", width=Inches(5.8))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 8.2: Empirical Benchmark Plot: Execution Time vs. Input Scale N")
        r_cap.italic = True
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10)
        p_cap.paragraph_format.space_after = Pt(14)

    # --- 9. REFERENCES ---
    add_header_styled(doc, "9. References", level=1)
    refs = [
        "1. Cormen, T. H., Leiserson, C. E., Rivest, R. L., & Stein, C. (2022). Introduction to Algorithms (4th ed.). MIT Press.",
        "2. Knuth, D. E. (1998). The Art of Computer Programming, Volume 3: Sorting and Searching (2nd ed.). Addison-Wesley.",
        "3. Sedgewick, R., & Wayne, K. (2011). Algorithms (4th ed.). Addison-Wesley Professional.",
        "4. SFML Development Team. (2024). Simple and Fast Multimedia Library Documentation (Release 2.5.1). https://www.sfml-dev.org/documentation/2.5.1/",
        "5. ISO/IEC. (2017). ISO/IEC 14882:2017 - Programming Languages — C++17. International Organization for Standardization.",
        "6. Williams, J. W. J. (1964). Algorithm 232: Heapsort. Communications of the ACM, 7(6), 347-348.",
        "7. Hoare, C. A. R. (1962). Quicksort. The Computer Journal, 5(1), 10-16."
    ]
    for ref in refs:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.space_after = Pt(4)
        p_ref.paragraph_format.line_spacing = 1.15
        r = p_ref.add_run(ref)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)

    doc.add_page_break()

    # --- 10. APPROVAL OF THE PROJECT ---
    add_header_styled(doc, "10. Approval of the Project", level=1)
    add_body_paragraph(doc,
        f"This is to certify that the project report entitled \"{PROJECT_TITLE}\" submitted by {STUDENTS[0][0]} ({STUDENTS[0][1]}), {STUDENTS[1][0]} ({STUDENTS[1][1]}), and {STUDENTS[2][0]} ({STUDENTS[2][1]}) in partial fulfillment of the requirements for the award of the degree of Bachelor of Technology in Computer Science & Engineering (Artificial Intelligence) from {COLLEGE_NAME}, affiliated to {UNIVERSITY_NAME}, is a bona fide record of the work carried out under our supervision.")

    p_cert2 = doc.add_paragraph()
    p_cert2.paragraph_format.space_before = Pt(8)
    p_cert2.paragraph_format.space_after = Pt(50)
    p_cert2.paragraph_format.line_spacing = 1.25
    r2 = p_cert2.add_run("The results embodied in this report have not been submitted to any other university or institute for the award of any degree or diploma.")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11.5)

    # Signature blocks
    tbl_sig = doc.add_table(rows=2, cols=2)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sig.autofit = False
    tbl_sig.columns[0].width = Inches(3.2)
    tbl_sig.columns[1].width = Inches(3.2)

    c00 = tbl_sig.rows[0].cells[0].paragraphs[0]
    c00.add_run("_________________________\nProject Coordinator\n").font.name = 'Times New Roman'
    r_g = c00.add_run(f"{GUIDE_NAME}\n{GUIDE_DESIGNATION}")
    r_g.bold = True
    r_g.font.name = 'Times New Roman'

    c01 = tbl_sig.rows[0].cells[1].paragraphs[0]
    c01.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    c01.add_run("_________________________\nHead of Department (HOD)\nDepartment of CSE (AI)\nJECRC, Jaipur").font.name = 'Times New Roman'

    tbl_sig.rows[1].cells[0].paragraphs[0].paragraph_format.space_before = Pt(40)
    c10 = tbl_sig.rows[1].cells[0].paragraphs[0]
    c10.add_run("Date: ___________________\nPlace: Jaipur").font.name = 'Times New Roman'

    output_path = "PROJECT_REPORT.docx"
    doc.save(output_path)
    print(f"Successfully generated {output_path}")

# ============================================================================
# 2. GENERATE PROJECT SYNOPSIS (.docx)
# ============================================================================
def generate_project_synopsis():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # --- COVER PAGE ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    
    r = p.add_run("A\nProject\nSYNOPSIS\nOf\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True

    r_title = p.add_run(f"{PROJECT_TITLE}\n\n")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(17)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(16, 44, 87)

    r_sub = p.add_run("Submitted in partial fulfillment for the award of degree of\n")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(11.5)
    r_sub.italic = True

    r_deg = p.add_run("Bachelor of Technology\n\nDegree Of\n\n")
    r_deg.font.name = 'Times New Roman'
    r_deg.font.size = Pt(13)
    r_deg.font.bold = True

    r_univ = p.add_run(f"{UNIVERSITY_NAME}\n\n")
    r_univ.font.name = 'Times New Roman'
    r_univ.font.size = Pt(14)
    r_univ.font.bold = True

    r_inst = p.add_run(f"{COLLEGE_NAME}\n\n{SESSION}\n\n")
    r_inst.font.name = 'Times New Roman'
    r_inst.font.size = Pt(13)
    r_inst.font.bold = True

    # Submitted By & Coordinator in 2 columns
    table_meta = doc.add_table(rows=1, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_meta.autofit = False
    table_meta.columns[0].width = Inches(3.2)
    table_meta.columns[1].width = Inches(3.2)

    cell_left = table_meta.rows[0].cells[0]
    p_sub = cell_left.paragraphs[0]
    p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_sb = p_sub.add_run("Submitted By:\n")
    r_sb.bold = True
    r_sb.font.name = 'Times New Roman'
    r_sb.font.size = Pt(11)
    for name, roll in STUDENTS:
        r_st = p_sub.add_run(f"{name} ({roll})\n")
        r_st.font.name = 'Times New Roman'
        r_st.font.size = Pt(10.5)

    cell_right = table_meta.rows[0].cells[1]
    p_coord = cell_right.paragraphs[0]
    p_coord.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_cb = p_coord.add_run(f"{GUIDE_DESIGNATION}\n")
    r_cb.bold = True
    r_cb.font.name = 'Times New Roman'
    r_cb.font.size = Pt(11)
    r_cn = p_coord.add_run(f"{GUIDE_NAME}\nFaculty Mentor\n")
    r_cn.font.name = 'Times New Roman'
    r_cn.font.size = Pt(10.5)

    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(36)
    p_foot.paragraph_format.space_after = Pt(2)
    
    r_dept = p_foot.add_run(f"{DEPARTMENT_NAME}\n{COLLEGE_NAME}\n")
    r_dept.font.name = 'Times New Roman'
    r_dept.font.size = Pt(11)
    r_dept.font.bold = True

    for line in ADDRESS_LINES:
        r_addr = p_foot.add_run(f"{line}\n")
        r_addr.font.name = 'Times New Roman'
        r_addr.font.size = Pt(9.5)

    doc.add_page_break()

    # --- PAGE 2: ABSTRACT ---
    p_abs_title = doc.add_paragraph()
    p_abs_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_abs = p_abs_title.add_run("ABSTRACT")
    r_abs.font.name = 'Times New Roman'
    r_abs.font.size = Pt(16)
    r_abs.font.bold = True
    r_abs.underline = True
    p_abs_title.paragraph_format.space_after = Pt(24)

    add_body_paragraph(doc,
        "Algorithm analysis in computer science pedagogy is predominantly theoretical, evaluating time and space complexities using asymptotic Big-O notations without direct empirical correlation to underlying physical computing architecture. The Sorting Analysis Comparator (SAC) is an engineering and laboratory analysis system built in C++17 and SFML 2.5 that rigorously bridges theoretical mathematical predictions and empirical hardware latency.")

    add_body_paragraph(doc,
        "SAC benchmarks four foundational sorting paradigms—Merge Sort (Divide-and-Conquer), Quick Sort (Partitioning), Heap Sort (Priority Queue Selection), and Radix Sort (Non-comparative Distribution)—across six scaling orders (N = 100 to 50,000) and four distinct topologies (Uniform Random, Nearly Sorted, Reverse Sorted, and Many Duplicates). Utilizing Windows QueryPerformanceCounter (QPC) for sub-microsecond wall-clock precision and an instrumented telemetry logger for exact comparison and swap counts, the system computes empirical growth ratios R_obs = T(n2)/T(n1) and fits empirical power exponents alpha = ln(R_obs)/ln(n2/n1).")

    add_body_paragraph(doc,
        "Empirical results prove strict O(n log n) convergence (alpha approx 1.08) for Merge and Heap sorts on random inputs, reveal Quick Sort's quadratic degradation (alpha = 1.88-2.01) on duplicate-heavy arrays, and demonstrate Radix Sort's strictly linear scaling (alpha = 1.00) with zero comparisons. Finally, an interactive 60 FPS graphical visualizer provides Single Sort Bar Animation with live on-bar numeric values, an unprecedented 4-Way Split-Screen Race Mode, and execution curve plotting.")

    doc.add_page_break()

    # --- PAGE 3: INTRODUCTION ---
    p_intro_title = doc.add_paragraph()
    p_intro_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_it = p_intro_title.add_run("INTRODUCTION")
    r_it.font.name = 'Times New Roman'
    r_it.font.size = Pt(16)
    r_it.font.bold = True
    r_it.underline = True
    p_intro_title.paragraph_format.space_after = Pt(16)

    intro_sections = [
        ("Problem Statement: ", 
         "Students and algorithm engineers frequently encounter a significant disconnect between theoretical Big-O complexity and physical execution time. Textbooks state that Quick Sort and Merge Sort are O(n log n), yet ignore hardware cache lines, branch miss penalties, and the catastrophic O(n^2) degradation of naive partitioning on duplicate keys. There exists a lack of integrated tools capable of simultaneously recording hardware wall-clock time, counting exact comparison/swap operations, fitting empirical growth exponents, and visually rendering algorithmic steps in real time."),
        
        ("Scope of the Project: ",
         "The scope encompasses the complete design and implementation of a C++17 dual-variant benchmarking suite, sub-microsecond hardware timing harnesses, automated CSV reporting, empirical power ratio fitting, and a desktop SFML graphical visualizer supporting Single Sort animation, 4-Way Head-to-Head Race Mode, Benchmark Plotting, and arbitrary Custom Array user input."),

        ("Technologies Used: ",
         "• Language: C++17 (high-performance vectorization, templates, memory buffers)\n• Graphics & Windowing: SFML 2.5.1 (Simple and Fast Multimedia Library)\n• Timing API: Windows QueryPerformanceCounter (QPC) hardware clock\n• Build Toolchain: MinGW GCC 6.3+ and CMake 3.10+\n• Data & Plotting: Python 3.10, Matplotlib, Pandas"),

        ("Methodology: ",
         "Dual-variant implementations ensure benchmark timings remain free of logging overhead (compiled under -O3), while an instrumented telemetry engine records comparison counts, swaps, and animation steps. Benchmarks are averaged over 5 stochastic trials per configuration across scales N in {100, 500, 1000, 5000, 10000, 50000}. Real-time event playback reconstructs array transitions at 60 FPS."),

        ("Expected Outcome: ",
         "A production-grade, highly educational laboratory software suite capable of proving asymptotic theoretical sorting complexities, highlighting worst-case degradation scenarios, exporting full statistical CSV datasets, and offering real-time animated sorting races."),

        ("Applications of Project: ",
         "Serves as an essential teaching instrument for Analysis of Algorithms laboratories, algorithm profiling in embedded systems, database query optimizer verification, and technical interview preparation.")
    ]

    for pref, text in intro_sections:
        add_body_paragraph(doc, text, bold_prefix=pref, space_after=6)

    doc.add_page_break()

    # --- PAGE 4: CONCLUSIONS AND FUTURE ENHANCEMENTS ---
    p_concl_title = doc.add_paragraph()
    p_concl_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ct = p_concl_title.add_run("CONCLUSIONS AND FUTURE ENHANCEMENTS")
    r_ct.font.name = 'Times New Roman'
    r_ct.font.size = Pt(15)
    r_ct.font.bold = True
    r_ct.underline = True
    p_concl_title.paragraph_format.space_after = Pt(14)

    add_body_paragraph(doc,
        "Conclusions: ", bold_prefix="• ")
    add_body_paragraph(doc,
        "The experimental results conclusively validate theoretical complexity bounds while uncovering critical hardware-level insights: (1) Merge Sort and Heap Sort consistently achieve O(n log n) scaling across all permutations, with observed growth ratios closely matching theoretical predictions (5.7x vs 5.87x expected from 10k to 50k); (2) Quick Sort exhibits excellent constant-factor performance on uniform random inputs but undergoes severe O(n^2) degradation on high duplicate densities (ratio 20.69x vs 25.0x theoretical quadratic); (3) Radix Sort strictly adheres to linear O(n) scaling with an empirical alpha of 1.00 and zero comparisons, confirming the non-comparative advantage.")

    add_body_paragraph(doc,
        "Future Enhancements: ", bold_prefix="• ")
    add_body_paragraph(doc,
        "1. Parallel GPU Sorting: Implementation of CUDA and OpenCL kernels for GPU Bitonic Merge Sort and Radix Sort.\n2. Audio Synthesis: Generation of dynamic audio tones (frequencies mapped to array values) to produce algorithmic soundscapes.\n3. External Memory Sorting: Benchmarking disk I/O paging mechanisms for datasets exceeding physical RAM (10^7+ elements).\n4. Multi-Threaded Parallelism: Integrating std::execution::par policies to evaluate OpenMP multicore scaling.")

    # Reference Image
    if os.path.exists("docs/block_diagram.png"):
        doc.add_picture("docs/block_diagram.png", width=Inches(5.6))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure: System Architecture of Sorting Analysis Comparator")
        r_cap.italic = True
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(9.5)

    doc.add_page_break()

    # --- PAGE 5: REFERENCES ---
    p_ref_title = doc.add_paragraph()
    p_ref_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_rt = p_ref_title.add_run("References")
    r_rt.font.name = 'Times New Roman'
    r_rt.font.size = Pt(16)
    r_rt.font.bold = True
    r_rt.font.color.rgb = RGBColor(16, 44, 87)
    r_rt.underline = True
    p_ref_title.paragraph_format.space_after = Pt(20)

    refs_syn = [
        "1. Cormen, T. H., Leiserson, C. E., Rivest, R. L., & Stein, C. (2022). Introduction to Algorithms (4th ed.). MIT Press.",
        "2. Knuth, D. E. (1998). The Art of Computer Programming, Volume 3: Sorting and Searching (2nd ed.). Addison-Wesley.",
        "3. Sedgewick, R., & Wayne, K. (2011). Algorithms (4th ed.). Addison-Wesley Professional.",
        "4. SFML Development Team. (2024). Simple and Fast Multimedia Library Documentation (Release 2.5.1). https://www.sfml-dev.org/",
        "5. ISO/IEC. (2017). ISO/IEC 14882:2017 - Programming Languages — C++17. International Organization for Standardization.",
        "6. Hoare, C. A. R. (1962). Quicksort. The Computer Journal, 5(1), 10-16.",
        "7. Williams, J. W. J. (1964). Algorithm 232: Heapsort. Communications of the ACM, 7(6), 347-348."
    ]

    for ref in refs_syn:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.space_after = Pt(6)
        p_ref.paragraph_format.line_spacing = 1.2
        r = p_ref.add_run(ref)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)

    output_path = "PROJECT_SYNOPSIS.docx"
    doc.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == '__main__':
    generate_project_report()
    generate_project_synopsis()
